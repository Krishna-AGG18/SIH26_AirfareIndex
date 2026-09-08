from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
from collections import Counter
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.config import get_settings
from app.db.models import AirfareIndexObservation
from app.db.urls import async_database_url

MONTHS = {
    "January": 1,
    "February": 2,
    "March": 3,
    "April": 4,
    "May": 5,
    "June": 6,
    "July": 7,
    "August": 8,
    "September": 9,
    "October": 10,
    "November": 11,
    "December": 12,
}
EXPECTED_HEADERS = {
    "base_year",
    "series",
    "year",
    "month",
    "state",
    "sector",
    "division",
    "group",
    "class",
    "sub_class",
    "item",
    "code",
    "index",
    "inflation",
    "imputation",
}
DATASET_KEY = "mospi-cpi-airfare-2025-2026"
SOURCE_NAME = "MoSPI CPI Airfare Index"
UPSERT_CHUNK_SIZE = 500


def load_rows(path: Path) -> tuple[list[dict[str, object]], str]:
    source_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != ["CPI Data"]:
        raise ValueError(f"Expected the CPI Data sheet, found {workbook.sheetnames}")

    worksheet = workbook["CPI Data"]
    values = list(worksheet.iter_rows(values_only=True))
    headers = set(values[0]) if values else set()
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"Unexpected headers: {sorted(headers)}")

    rows: list[dict[str, object]] = []
    for row_number, values_row in enumerate(values[1:], start=2):
        row = dict(zip(values[0], values_row, strict=True))
        if row["item"] != "Airfare" or row["code"] != "07.3.3.1.2.01":
            raise ValueError(f"Row {row_number} is not the expected airfare series")
        if row["month"] not in MONTHS:
            raise ValueError(f"Row {row_number} has an unknown month: {row['month']}")
        if row["index"] in (None, ""):
            raise ValueError(f"Row {row_number} has no index value")

        rows.append(
            {
                "dataset_key": DATASET_KEY,
                "source_name": SOURCE_NAME,
                "source_file": path.name,
                "source_sha256": source_sha256,
                "base_year": int(row["base_year"]),
                "series": str(row["series"]),
                "observation_month": date(int(row["year"]), MONTHS[str(row["month"])], 1),
                "state": str(row["state"]),
                "sector": str(row["sector"]),
                "division": str(row["division"]),
                "group_name": str(row["group"]),
                "class_name": str(row["class"]),
                "sub_class": str(row["sub_class"]),
                "item": str(row["item"]),
                "code": str(row["code"]),
                "index_value": float(row["index"]),
                "inflation_yoy": float(row["inflation"]) if row["inflation"] not in (None, "") else None,
                "imputation": str(row["imputation"]).upper() == "Y",
            }
        )

    workbook.close()
    return rows, source_sha256


async def import_rows(rows: list[dict[str, object]]) -> None:
    settings = get_settings()
    if not settings.database_url_unpooled:
        raise RuntimeError("DATABASE_URL_UNPOOLED is required for this import")

    engine = create_async_engine(async_database_url(settings.database_url_unpooled), pool_pre_ping=True)
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    conflict_key = ["dataset_key", "observation_month", "state", "sector", "code"]

    async with session_factory() as session:
        for start in range(0, len(rows), UPSERT_CHUNK_SIZE):
            statement = insert(AirfareIndexObservation).values(rows[start : start + UPSERT_CHUNK_SIZE])
            update_columns = {
                column.name: getattr(statement.excluded, column.name)
                for column in AirfareIndexObservation.__table__.columns
                if column.name not in {"id", *conflict_key, "imported_at"}
            }
            await session.execute(
                statement.on_conflict_do_update(
                    index_elements=conflict_key,
                    set_=update_columns,
                )
            )
        await session.commit()

    await engine.dispose()


def main() -> None:
    parser = argparse.ArgumentParser(description="Import MoSPI CPI airfare index rows into Neon synthetic data")
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    rows, source_sha256 = load_rows(args.workbook)
    asyncio.run(import_rows(rows))

    periods = sorted({row["observation_month"].isoformat() for row in rows})
    print(
        json.dumps(
            {
                "dataset_key": DATASET_KEY,
                "source_file": args.workbook.name,
                "source_sha256": source_sha256,
                "rows_upserted": len(rows),
                "periods": {"first": periods[0], "last": periods[-1], "count": len(periods)},
                "states": len({row["state"] for row in rows}),
                "sectors": Counter(row["sector"] for row in rows),
                "imputed_rows": sum(bool(row["imputation"]) for row in rows),
                "imported_at_utc": datetime.now(UTC).isoformat(timespec="seconds"),
            },
            indent=2,
            default=str,
        )
    )


if __name__ == "__main__":
    main()
